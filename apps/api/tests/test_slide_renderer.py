import io

import pymupdf
from pptx import Presentation

from app.rendering.slide_renderer import hash_slide_xml, rasterize_page_to_png


def _one_slide_deck(title: str) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text_frame.text = title
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_hash_slide_xml_is_deterministic():
    deck = _one_slide_deck("Hello")
    presentation = Presentation(io.BytesIO(deck))
    slide = presentation.slides[0]

    assert hash_slide_xml(slide) == hash_slide_xml(slide)


def test_hash_slide_xml_changes_with_content():
    hash_a = hash_slide_xml(Presentation(io.BytesIO(_one_slide_deck("Hello"))).slides[0])
    hash_b = hash_slide_xml(Presentation(io.BytesIO(_one_slide_deck("Goodbye"))).slides[0])

    assert hash_a != hash_b


def _one_page_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Rendered slide test")
    data = doc.tobytes()
    doc.close()
    return data


def test_rasterize_page_to_png_produces_valid_png_at_expected_dpi():
    png_bytes = rasterize_page_to_png(_one_page_pdf(), page_index=0, dpi=150)

    assert png_bytes.startswith(b"\x89PNG\r\n\x1a\n")

    pixmap = pymupdf.Pixmap(png_bytes)
    # A default new_page() is A4 (595x842pt); at 150 DPI that's ~1240x1754px.
    assert abs(pixmap.width - 1240) < 5
    assert abs(pixmap.height - 1754) < 5


def test_resolve_soffice_path_returns_none_when_unavailable(monkeypatch):
    from app.rendering.slide_renderer import _resolve_soffice_path
    monkeypatch.setattr("shutil.which", lambda _: None)
    monkeypatch.setattr("pathlib.Path.is_file", lambda _: False)
    assert _resolve_soffice_path() is None


async def test_convert_pptx_to_pdf_raises_unavailable_when_neither_available(monkeypatch):
    import pytest
    from app.rendering.slide_renderer import convert_pptx_to_pdf_bytes
    from app.rendering.subprocess_runner import SubprocessUnavailableError

    monkeypatch.setattr("app.rendering.slide_renderer._resolve_soffice_path", lambda: None)
    monkeypatch.setattr("app.rendering.slide_renderer._is_powerpoint_available", lambda: False)

    with pytest.raises(SubprocessUnavailableError) as exc_info:
        await convert_pptx_to_pdf_bytes(b"dummy")
    assert "No slide renderer available" in str(exc_info.value)
