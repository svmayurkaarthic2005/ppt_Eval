import io

import pytest
from pptx import Presentation

from app.extraction.slide_roles import classify_by_keywords
from app.extraction.structural import PptxCorruptError, parse_numeric_cell, parse_pptx
from app.models.schemas import SlideRole


def _build_deck(titles: list[str]) -> bytes:
    """A deck with len(titles) slides, each with a title and one body line."""
    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # "Title and Content"
    for index, title in enumerate(titles, start=1):
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text_frame.text = title
        slide.placeholders[1].text_frame.text = f"Body text for slide {index}."
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_parses_5_slide_deck_without_hardcoding_slide_count():
    deck = _build_deck([f"Slide {i}" for i in range(1, 6)])
    slides = parse_pptx(deck)

    assert len(slides) == 5
    assert [s.slide_number for s in slides] == [1, 2, 3, 4, 5]
    assert slides[0].title == "Slide 1"
    assert "Body text for slide 1." in slides[0].raw_text


def test_parses_17_slide_deck_without_hardcoding_slide_count():
    deck = _build_deck([f"Slide {i}" for i in range(1, 18)])
    slides = parse_pptx(deck)

    assert len(slides) == 17
    assert [s.slide_number for s in slides] == list(range(1, 18))


def test_title_is_kept_separate_from_body_text():
    deck = _build_deck(["The Problem We're Solving"])
    slides = parse_pptx(deck)

    assert slides[0].title == "The Problem We're Solving"
    assert "The Problem We're Solving" not in slides[0].raw_text


def test_corrupt_pptx_raises_pptx_corrupt_error():
    with pytest.raises(PptxCorruptError):
        parse_pptx(b"this is not a valid zip/pptx file at all")


def test_table_numeric_cells_parsed_without_coercing_malformed_to_zero():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(2, 2, 0, 0, 1000, 1000)
    table = graphic_frame.table
    table.cell(0, 0).text = "$1,200.50"
    table.cell(0, 1).text = "N/A"
    table.cell(1, 0).text = "-300"
    table.cell(1, 1).text = "42"
    buffer = io.BytesIO()
    presentation.save(buffer)

    slides = parse_pptx(buffer.getvalue())
    rows = slides[0].tables[0]

    assert rows[0][0] == {"text": "$1,200.50", "value": 1200.5}
    assert rows[0][1] == {"text": "N/A", "value": None}
    assert rows[1][0] == {"text": "-300", "value": -300}
    assert rows[1][1] == {"text": "42", "value": 42}


def test_parse_numeric_cell():
    assert parse_numeric_cell("$1,200.50") == 1200.5
    assert parse_numeric_cell("1,000") == 1000
    assert parse_numeric_cell("N/A") is None
    assert parse_numeric_cell("") is None
    assert parse_numeric_cell("-42") == -42


def test_hyperlink_extraction():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(0, 0, 1000, 1000)
    paragraph = textbox.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "our repo"
    run.hyperlink.address = "https://github.com/example/repo"
    buffer = io.BytesIO()
    presentation.save(buffer)

    slides = parse_pptx(buffer.getvalue())
    shape = next(s for s in slides[0].shapes if s["hyperlinks"])
    assert shape["hyperlinks"] == [
        {"text": "our repo", "address": "https://github.com/example/repo"}
    ]


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("The Problem", SlideRole.problem),
        ("Existing Solutions in the Market", SlideRole.existing_solution),
        ("Our Proposed Solution", SlideRole.proposed_solution),
        ("System Architecture", SlideRole.workflow),
        ("Pricing & Cost Breakdown", SlideRole.cost),
        ("Project Timeline", SlideRole.timeline),
    ],
)
def test_classify_by_keywords_matches_expected_role(title, expected):
    assert classify_by_keywords(title) == expected


def test_classify_by_keywords_returns_none_when_ambiguous_or_no_match():
    assert classify_by_keywords("Thank You!") is None
    assert classify_by_keywords("") is None
